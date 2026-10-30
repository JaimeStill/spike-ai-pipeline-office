# /// script
# requires-python = ">=3.12"
# dependencies = ["openpyxl==3.1.5", "python-docx==1.2.0", "python-pptx==1.0.2"]
# ///
"""Generate the office fixtures and their ground truth.

Each fixture is written twice: by a library from the definitions below
(<name>.python-docx.docx, <name>.openpyxl.xlsx, or <name>.python-pptx.pptx), and
that file re-saved by LibreOffice headless (<name>.libreoffice.docx, .xlsx, or
.pptx), so the corpus carries
a real application's XML as well as a library's. Both files of a fixture share
one ground truth, <name>.truth.json, which this script writes from the same
definitions, unless LibreOffice changes what the document holds, when the
re-save has its own, <name>.libreoffice.truth.json. The ground truth is the
document's text in its accepted state, element by element, in a shape of its
own, so it never depends on the office model it checks.

python-docx writes no tracked changes, so the w:ins, w:del, w:moveFrom, and
w:moveTo elements and the deleted paragraph marks are written as raw XML;
openpyxl writes no shared formulas, no ISO-date cells, and no pictures'
alternative text, so those are patched in as raw XML too; python-pptx has no API
for hidden slides, slide order, a paragraph's own bullet, or pictures'
alternative text, so those are set on its XML.

Run it through `mise run fixtures`; the check never runs it. It rewrites the
generated set and the checksum section of fixtures/README.md.
"""

import datetime
import hashlib
import io
import json
import pathlib
import shutil
import struct
import subprocess
import sys
import tempfile
import zipfile
import zlib
from contextlib import contextmanager
from dataclasses import dataclass
from xml.sax.saxutils import escape

import docx
import openpyxl
import pptx
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.worksheet.formula import ArrayFormula
from docx.enum.section import WD_SECTION
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.parts.image import ImagePart
from docx.shared import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

ROOT = pathlib.Path(__file__).resolve().parent
GENERATED = ROOT / "generated"
README = ROOT / "README.md"
AUTHOR = "Fixture Author"
DATE = "2026-01-15T10:00:00Z"
W = nsdecls("w")
NS = nsdecls("w", "wp", "a", "pic", "r", "c")
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
CT_COMMENTS_EXTENDED = "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml"
RT_COMMENTS_EXTENDED = "http://schemas.microsoft.com/office/2011/relationships/commentsExtended"
EMU = 457200  # A drawing's side, half an inch.


class Fixture:
    """A python-docx document and its ground truth, built side by side."""

    writer, extension, filter = "python-docx", "docx", "docx:MS Word 2007 XML"

    def __init__(self, name: str):
        self.name = name
        self.doc = docx.Document()
        self.elements: list[dict] = []
        self.absent: list[str] = []
        self.once: list[str] = []
        self._ids = 100
        self._numbering_ready = False
        # The story blocks are written to (the body, or a header or footer) and the section index
        # the blocks are in.
        self.target = self.doc
        self.section = 0
        # Each section's headers and footers in the ground truth, when the fixture declares them.
        self.sections: list[dict] | None = None
        self.notes: dict[str, dict] = {"footnote": {}, "endnote": {}}
        self._note_order: dict[str, list[str]] = {"footnote": [], "endnote": []}
        self.note_truth: dict[str, list[dict]] = {"footnote": [], "endnote": []}
        self.comments: dict[str, dict] = {}
        self.comments_extended = True
        # How the LibreOffice re-save's ground truth differs, when it does: a function that
        # adjusts a copy of the truth, written as <name>.libreoffice.truth.json.
        self.libreoffice = None
        # Comments python-docx writes itself, with their ground truth.
        self.plain_comments: list[dict] = []
        self._raw_parts: dict[str, ImagePart] = {}
        self._charts = 0

    def next_id(self) -> int:
        self._ids += 1
        return self._ids

    # Plain blocks.

    def heading(self, text: str, level: int) -> None:
        self.doc.add_heading(text, level)
        self.elements.append({"kind": "heading", "level": level, "text": text})

    def outline_heading(self, text: str, level: int) -> None:
        """A Normal-style paragraph made a heading by its outline level alone."""
        p = self.target.add_paragraph(text)
        p._p.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {W} w:val="{level - 1}"/>'))
        self.elements.append({"kind": "heading", "level": level, "text": text})

    def paragraph(self, text: str) -> None:
        self.target.add_paragraph(text)
        self.elements.append({"kind": "paragraph", "text": text})

    def paragraph_with_break(self, first: str, second: str) -> None:
        p = self.target.add_paragraph()
        run = p.add_run(first)
        run.add_break()
        p.add_run(second)
        self.elements.append({"kind": "paragraph", "text": f"{first}\n{second}"})

    # Lists, numbered through definitions this script adds to the numbering part.

    def _numbering(self) -> None:
        if self._numbering_ready:
            return
        numbering = self.doc.part.numbering_part.element

        def lvl(ilvl: int, fmt: str, text: str) -> str:
            indent = 720 * (ilvl + 1)
            return (
                f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
                f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
                f'<w:pPr><w:ind w:left="{indent}" w:hanging="360"/></w:pPr></w:lvl>'
            )

        bullets = lvl(0, "bullet", "•") + lvl(1, "bullet", "◦") + lvl(2, "bullet", "▪")
        mixed = lvl(0, "decimal", "%1.") + lvl(1, "bullet", "•") + lvl(2, "lowerLetter", "%3)")
        abstracts = [
            f'<w:abstractNum {W} w:abstractNumId="10"><w:multiLevelType w:val="hybridMultilevel"/>{bullets}</w:abstractNum>',
            f'<w:abstractNum {W} w:abstractNumId="11"><w:multiLevelType w:val="hybridMultilevel"/>{mixed}</w:abstractNum>',
        ]
        nums = [
            f'<w:num {W} w:numId="10"><w:abstractNumId w:val="10"/></w:num>',
            f'<w:num {W} w:numId="11"><w:abstractNumId w:val="11"/></w:num>',
        ] + [
            # Instances of one abstract definition share its count unless they restart it.
            f'<w:num {W} w:numId="{num_id}"><w:abstractNumId w:val="11"/>'
            f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'
            for num_id in (12, 13, 14)
        ]
        # The schema orders every w:abstractNum before every w:num.
        first_num = numbering.find(f"{{{numbering.nsmap['w']}}}num")
        for xml in abstracts:
            el = parse_xml(xml)
            if first_num is None:
                numbering.append(el)
            else:
                first_num.addprevious(el)
        for xml in nums:
            numbering.append(parse_xml(xml))
        self._numbering_ready = True

    def _number(self, paragraph, num_id: int, level: int) -> None:
        self._numbering()
        num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = level
        num_pr.get_or_add_numId().val = num_id

    def _list_item(self, text: str, level: int, ordered: bool, new_list: bool) -> None:
        """The ground truth of a list item. It starts a list when it follows anything but a
        list item, when it is deeper than the item before it, or when the definition says it
        starts a new list at its level (new_list)."""
        previous = self.elements[-1] if self.elements else None
        start = new_list or previous is None or previous["kind"] != "listItem" or previous["level"] < level
        self.elements.append(
            {"kind": "listItem", "level": level, "ordered": ordered, "start": start, "text": text}
        )

    def item(self, text: str, num_id: int, level: int, ordered: bool, new_list: bool = False) -> None:
        p = self.target.add_paragraph(text, style="List Paragraph")
        self._number(p, num_id, level)
        self._list_item(text, level, ordered, new_list)

    # Tables. Each cell is (text, colSpan, rowSpan); None marks a covered grid position.

    def table(self, grid: list[list[tuple[str, int, int] | None]]) -> None:
        rows, cols = len(grid), len(grid[0])
        table = self._add_table(rows, cols)
        table.style = "Table Grid"
        truth_rows = []
        for r, row in enumerate(grid):
            truth_row = []
            for c, spec in enumerate(row):
                if spec is None:
                    continue
                text, col_span, row_span = spec
                cell = table.cell(r, c)
                if col_span > 1 or row_span > 1:
                    cell = cell.merge(table.cell(r + row_span - 1, c + col_span - 1))
                paragraphs = text.split("\n")
                cell.paragraphs[0].text = paragraphs[0]
                for extra in paragraphs[1:]:
                    cell.add_paragraph(extra)
                truth_row.append({"text": text, "colSpan": col_span, "rowSpan": row_span})
            truth_rows.append(truth_row)
        self.elements.append({"kind": "table", "rows": truth_rows})

    # Segments, as raw XML. A segment is (kind, value):
    # - "text", "ins", "del", "moveFrom", "moveTo": text, tracked or not (one move per fixture,
    #   named move1); the accepted text keeps "text", "ins", and "moveTo";
    # - "footnote", "endnote": the key of a note this fixture defines, referenced here;
    #   "delNote": (kind, key), a reference inside a deletion, so the note is never referenced;
    # - "commentStart", "commentEnd": the key of a comment this fixture defines, whose anchored
    #   range (and its replies') starts or ends here;
    # - "picture", "delPicture": a Picture, in a drawing of its own, or inside a deletion;
    # - "chart": a Chart, in a drawing of its own.
    # References, pictures, and charts add no text; a reference marks its offset in the text, and
    # a picture or chart is a block of its own after its paragraph.

    def _run(self, text: str, deleted: bool = False) -> str:
        tag = "w:delText" if deleted else "w:t"
        return f'<w:r><{tag} xml:space="preserve">{escape(text)}</{tag}></w:r>'

    def _segments(self, segments: list[tuple[str, object]], part) -> tuple[str, list[dict], list[dict]]:
        """The XML of segments in the given story part, the marks of the references they hold,
        and the blocks (media and charts) that follow their paragraph."""
        attrs = f'w:author="{AUTHOR}" w:date="{DATE}"'
        xml, marks, after = [], [], []
        offset = 0
        for kind, value in segments:
            if kind == "text":
                xml.append(self._run(value))
            elif kind == "ins":
                xml.append(f'<w:ins w:id="{self.next_id()}" {attrs}>{self._run(value)}</w:ins>')
            elif kind == "del":
                xml.append(f'<w:del w:id="{self.next_id()}" {attrs}>{self._run(value, True)}</w:del>')
            elif kind in ("moveFrom", "moveTo"):
                range_id = self.next_id()
                xml.append(
                    f'<w:{kind}RangeStart w:id="{range_id}" {attrs} w:name="move1"/>'
                    f'<w:{kind} w:id="{self.next_id()}" {attrs}>{self._run(value)}</w:{kind}>'
                    f'<w:{kind}RangeEnd w:id="{range_id}"/>'
                )
            elif kind in ("footnote", "endnote"):
                xml.append(self._note_reference(kind, value))
                marks.append({"at": offset, "kind": kind, "n": self._reference_note(kind, value)})
            elif kind == "delNote":
                note_kind, key = value
                xml.append(f'<w:del w:id="{self.next_id()}" {attrs}>{self._note_reference(note_kind, key)}</w:del>')
            elif kind == "commentStart":
                thread = self.comments[value]
                xml.append("".join(f'<w:commentRangeStart w:id="{i}"/>' for i in thread["ids"]))
            elif kind == "commentEnd":
                thread = self.comments[value]
                xml.append(
                    "".join(
                        f'<w:commentRangeEnd w:id="{i}"/>'
                        f'<w:r><w:rPr><w:rStyle w:val="CommentReference"/></w:rPr><w:commentReference w:id="{i}"/></w:r>'
                        for i in thread["ids"]
                    )
                )
                marks.append({"at": offset, "kind": "comment", "n": list(self.comments).index(value) + 1})
            elif kind == "picture":
                xml.append(f"<w:r>{self._picture_drawing(value, part)}</w:r>")
                after.append(value.truth(self.section))
            elif kind == "delPicture":
                xml.append(f'<w:del w:id="{self.next_id()}" {attrs}><w:r>{self._picture_drawing(value, part)}</w:r></w:del>')
            elif kind == "chart":
                xml.append(f"<w:r>{self._chart_drawing(value, part)}</w:r>")
                after.append(value.truth())
            else:
                raise ValueError(kind)
            if kind in ("text", "ins", "moveTo"):
                offset += len(value)
        return "".join(xml), marks, after

    @staticmethod
    def accepted(segments: list[tuple[str, object]]) -> str:
        return "".join(value for kind, value in segments if kind in ("text", "ins", "moveTo"))

    def _fill(self, p, segments) -> tuple[list[dict], list[dict]]:
        """Appends segments to a python-docx paragraph; returns their marks and the blocks after."""
        xml, marks, after = self._segments(segments, p.part)
        for el in parse_xml(f"<w:body {NS}>{xml}</w:body>"):
            p._p.append(el)
        return marks, after

    def _tracked_paragraph(self, segments, mark: str | None = None, style: str | None = None):
        """A paragraph of segments; mark "del" or "ins" tracks its paragraph mark. Returns the
        paragraph, the marks of its references, and the blocks after it."""
        p = self.target.add_paragraph(style=style)
        marks, after = self._fill(p, segments)
        if mark is not None:
            p._p.get_or_add_pPr().append(
                parse_xml(
                    f'<w:rPr {W}><w:{mark} w:id="{self.next_id()}" w:author="{AUTHOR}" w:date="{DATE}"/></w:rPr>'
                )
            )
        return p, marks, after

    def _paragraph_truth(self, text: str, marks: list[dict], after: list[dict]) -> None:
        """A paragraph's ground truth: the paragraph unless it holds no text and no reference
        (then it is spacing), then the blocks after it."""
        if text.strip() or marks:
            element = {"kind": "paragraph", "text": text}
            if marks:
                element["marks"] = marks
            self.elements.append(element)
        self.elements.extend(after)

    def tracked(self, segments, mark: str | None = None) -> None:
        """A paragraph of segments, its mark tracked or not."""
        _, marks, after = self._tracked_paragraph(segments, mark)
        if mark != "del":
            self._paragraph_truth(self.accepted(segments), marks, after)

    def tracked_joined(self, first, second) -> None:
        """Two paragraphs, the first's mark deleted: accepted, they are one."""
        _, marks1, after1 = self._tracked_paragraph(first, "del")
        _, marks2, after2 = self._tracked_paragraph(second)
        shift = len(self.accepted(first))
        marks = marks1 + [{**m, "at": m["at"] + shift} for m in marks2]
        self._paragraph_truth(self.accepted(first) + self.accepted(second), marks, after1 + after2)

    def tracked_before_table(self, segments) -> None:
        """A paragraph whose mark was deleted, right before a table: accepted, its text stays a
        paragraph of its own before the table, as nothing follows it to join."""
        _, marks, after = self._tracked_paragraph(segments, "del")
        self._paragraph_truth(self.accepted(segments), marks, after)

    def tracked_item(self, segments, num_id: int, level: int, ordered: bool) -> None:
        p, marks, _ = self._tracked_paragraph(segments, style="List Paragraph")
        self._number(p, num_id, level)
        self._list_item(self.accepted(segments), level, ordered, False)
        if marks:
            self.elements[-1]["marks"] = marks

    def tracked_table(self, cells: list[list[tuple[str, object]]]) -> None:
        """A one-row table, a paragraph of segments in each cell. A cell's pictures and charts
        are listed as its media."""
        table = self._add_table(1, len(cells))
        table.style = "Table Grid"
        row = []
        for c, segments in enumerate(cells):
            _, after = self._fill(table.cell(0, c).paragraphs[0], segments)
            cell = {"text": self.accepted(segments), "colSpan": 1, "rowSpan": 1}
            if after:
                cell["media"] = after
            row.append(cell)
        self.elements.append({"kind": "table", "rows": [row]})

    def _add_table(self, rows: int, cols: int):
        # A header's or footer's table needs a width; the body's takes the page's.
        if self.target is self.doc:
            return self.doc.add_table(rows=rows, cols=cols)
        return self.target.add_table(rows, cols, Inches(6))

    # Stories other than the body: headers and footers, by python-docx; the section index their
    # media are placed in is the section they are read for.

    @contextmanager
    def story(self, container, section: int):
        """Builds into a header or footer; yields the ground truth of its blocks."""
        saved = self.target, self.elements, self.section
        self.target, self.elements, self.section = container, [], section
        try:
            yield self.elements
        finally:
            self.target, self.elements, self.section = saved

    def new_section(self):
        """Starts a new section; returns it, its header and footer references inherited."""
        self.section += 1
        return self.doc.add_section(WD_SECTION.NEW_PAGE)

    # Pictures and charts, in drawings written as raw XML.

    def _picture_rid(self, pic: "Picture", part) -> str:
        if pic.link is not None:
            return part.relate_to(pic.link, RT.IMAGE, is_external=True)
        if pic.part is not None:
            # A part of its own, so two parts can hold the same bytes; python-docx would share one.
            if pic.part not in self._raw_parts:
                self._raw_parts[pic.part] = ImagePart(PackURI(pic.part), pic.content_type, pic.data)
            return part.relate_to(self._raw_parts[pic.part], RT.IMAGE)
        rid, _ = part.get_or_add_image(io.BytesIO(pic.data))
        return rid

    def _picture_drawing(self, pic: "Picture", part) -> str:
        rid = self._picture_rid(pic, part)
        blip = f'r:link="{rid}"' if pic.link is not None else f'r:embed="{rid}"'
        shape_id = self.next_id()
        return (
            f"<w:drawing>{self._inline_open(shape_id, pic.descr, pic.title)}"
            f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            f'<pic:pic><pic:nvPicPr><pic:cNvPr id="{shape_id}" name="Picture {shape_id}"/><pic:cNvPicPr/></pic:nvPicPr>'
            f"<pic:blipFill><a:blip {blip}/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>"
            f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{EMU}" cy="{EMU}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic>'
            f"</a:graphicData></a:graphic></wp:inline></w:drawing>"
        )

    def _chart_drawing(self, chart: "Chart", part) -> str:
        self._charts += 1
        chart_part = Part(
            PackURI(f"/word/charts/chart{self._charts}.xml"), CT.DML_CHART, chart.xml().encode(), self.doc.part.package
        )
        rid = part.relate_to(chart_part, RT.CHART)
        shape_id = self.next_id()
        return (
            f'<w:drawing>{self._inline_open(shape_id, "", "")}'
            f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
            f'<c:chart r:id="{rid}"/></a:graphicData></a:graphic></wp:inline></w:drawing>'
        )

    @staticmethod
    def _inline_open(shape_id: int, descr: str, title: str) -> str:
        alt = (f' descr="{escape(descr)}"' if descr else "") + (f' title="{escape(title)}"' if title else "")
        return (
            f'<wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="{EMU}" cy="{EMU}"/>'
            f'<wp:docPr id="{shape_id}" name="Drawing {shape_id}"{alt}/>'
            f'<wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>'
        )

    # Footnotes and endnotes, written as raw parts with Word's separator and continuation notes.
    # A note's ground truth is listed in the order the body first references it.

    def note(self, kind: str, key: str, paragraphs: list[str]) -> None:
        """Defines a footnote or endnote, its ids numbered in definition order."""
        notes = self.notes[kind]
        notes[key] = {"id": len(notes) + 1, "paragraphs": paragraphs}

    def _note_reference(self, kind: str, key: str) -> str:
        return (
            f'<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
            f'<w:{kind}Reference w:id="{self.notes[kind][key]["id"]}"/></w:r>'
        )

    def _reference_note(self, kind: str, key: str) -> int:
        order = self._note_order[kind]
        if key not in order:
            order.append(key)
            # The first paragraph starts with the reference mark and Word's space after it.
            paragraphs = self.notes[kind][key]["paragraphs"]
            texts = [" " + paragraphs[0]] + paragraphs[1:]
            self.note_truth[kind].append({"elements": [{"kind": "paragraph", "text": t} for t in texts]})
        return order.index(key) + 1

    def _notes_part(self, kind: str) -> None:
        notes = self.notes[kind]
        if not notes:
            return
        tag = f"w:{kind}"
        separators = (
            f'<{tag} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></{tag}>'
            f'<{tag} w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></{tag}>'
        )
        body = []
        for note in notes.values():
            paragraphs = []
            for i, text in enumerate(note["paragraphs"]):
                ref = f"<w:r><w:rPr><w:vertAlign w:val=\"superscript\"/></w:rPr><w:{kind}Ref/></w:r>" if i == 0 else ""
                lead = " " if i == 0 else ""
                paragraphs.append(f"<w:p>{ref}{self._run(lead + text)}</w:p>")
            body.append(f'<{tag} w:id="{note["id"]}">{"".join(paragraphs)}</{tag}>')
        xml = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<{tag}s {W}>{separators}{"".join(body)}</{tag}s>'
        content_type, reltype = (CT.WML_FOOTNOTES, RT.FOOTNOTES) if kind == "footnote" else (CT.WML_ENDNOTES, RT.ENDNOTES)
        part = Part(PackURI(f"/word/{kind}s.xml"), content_type, xml.encode(), self.doc.part.package)
        self.doc.part.relate_to(part, reltype)

    # Comments, written as a raw comments part and, unless the fixture leaves it out, a
    # commentsExtended part giving each reply its parent and each thread its resolved state. A
    # thread's root and replies share its anchored range, as Word writes them.

    def comment(self, key: str, author: str, paragraphs: list[str], anchor: str, resolved: bool = False,
                replies: list[tuple[str, list[str]]] = ()) -> None:
        start = sum(len(t["ids"]) for t in self.comments.values())
        self.comments[key] = {
            "ids": list(range(start, start + 1 + len(replies))),
            "author": author,
            "paragraphs": paragraphs,
            "anchor": anchor,
            "resolved": resolved,
            "replies": list(replies),
        }

    def _comments_parts(self) -> None:
        if not self.comments:
            return
        comments, extended = [], []
        para = 0
        for thread in self.comments.values():
            entries = [(thread["author"], thread["paragraphs"])] + thread["replies"]
            root_para = None
            for comment_id, (author, texts) in zip(thread["ids"], entries):
                paragraphs = []
                for i, text in enumerate(texts):
                    para += 1
                    ref = "<w:r><w:annotationRef/></w:r>" if i == 0 else ""
                    paragraphs.append(f'<w:p w14:paraId="{para:08X}" w14:textId="77777777">{ref}{self._run(text)}</w:p>')
                initials = "".join(word[0] for word in author.split())
                comments.append(
                    f'<w:comment w:id="{comment_id}" w:author="{escape(author)}" w:date="{DATE}" w:initials="{initials}">'
                    f'{"".join(paragraphs)}</w:comment>'
                )
                parent = f' w15:paraIdParent="{root_para:08X}"' if root_para is not None else ""
                extended.append(f'<w15:commentEx w15:paraId="{para:08X}"{parent} w15:done="{int(thread["resolved"])}"/>')
                root_para = root_para if root_para is not None else para
        xml = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:comments {nsdecls("w", "w14")}>{"".join(comments)}</w:comments>'
        part = Part(PackURI("/word/comments.xml"), CT.WML_COMMENTS, xml.encode(), self.doc.part.package)
        self.doc.part.relate_to(part, RT.COMMENTS)
        if not self.comments_extended:
            return
        xml = (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            f'<w15:commentsEx xmlns:w15="{W15}">{"".join(extended)}</w15:commentsEx>'
        )
        part = Part(PackURI("/word/commentsExtended.xml"), CT_COMMENTS_EXTENDED, xml.encode(), self.doc.part.package)
        self.doc.part.relate_to(part, RT_COMMENTS_EXTENDED)

    def save(self, path: pathlib.Path) -> None:
        """Writes the parts the definitions add, notes and comments, then the document."""
        self._notes_part("footnote")
        self._notes_part("endnote")
        self._comments_parts()
        self.doc.save(path)

    def libreoffice_truth(self) -> dict | None:
        """The re-save's own ground truth, or None when it shares the fixture's."""
        if self.libreoffice is None:
            return None
        adjusted = json.loads(json.dumps(self.truth()))
        self.libreoffice(adjusted)
        return adjusted

    def truth(self) -> dict:
        truth = {
            "fixture": self.name,
            "format": "docx",
            "elements": self.elements,
            "absent": self.absent,
            "once": self.once,
        }
        if self.sections is not None:
            truth["sections"] = self.sections
        if self.note_truth["footnote"]:
            truth["footnotes"] = self.note_truth["footnote"]
        if self.note_truth["endnote"]:
            truth["endnotes"] = self.note_truth["endnote"]
        if self.comments or self.plain_comments:
            truth["comments"] = self.plain_comments + [
                {
                    "author": t["author"],
                    "anchor": t["anchor"],
                    "resolved": t["resolved"],
                    "elements": [{"kind": "paragraph", "text": p} for p in t["paragraphs"]],
                    "replies": [
                        {"author": a, "elements": [{"kind": "paragraph", "text": p} for p in ps]}
                        for a, ps in t["replies"]
                    ],
                }
                for t in self.comments.values()
            ]
        return truth


@dataclass
class Picture:
    """A picture: embedded, through python-docx's image parts (one per distinct content) or a
    part of its own, or linked to an external target, with no part."""

    descr: str
    title: str = ""
    data: bytes = b""
    ext: str = "png"
    content_type: str = CT.PNG
    part: str | None = None
    link: str | None = None

    def truth(self, section: int) -> dict:
        element = {"kind": "media"}
        if self.link is None:
            element["name"] = hashlib.sha256(self.data).hexdigest() + "." + self.ext
            element["contentType"] = self.content_type
        else:
            element["target"] = self.link
        if self.descr:
            element["alt"] = self.descr
        if self.title:
            element["title"] = self.title
        element["section"] = section
        return element


@dataclass
class Chart:
    """A chart part's data: its type (the plot element's name, less "Chart"), title, and series
    of name, categories, and values, a missing value None. The series are cached in the chart
    part, with no embedded workbook."""

    chart_type: str
    title: str | None
    series: list[tuple[str, list[str], list[float | None]]]

    def truth(self) -> dict:
        element = {"kind": "chart", "chartType": self.chart_type}
        if self.title is not None:
            element["title"] = self.title
        element["series"] = [{"name": n, "categories": c, "values": v} for n, c, v in self.series]
        return element

    def xml(self) -> str:
        def cache(tag: str, values: list) -> str:
            points = "".join(
                f'<c:pt idx="{i}"><c:v>{escape(str(v))}</c:v></c:pt>' for i, v in enumerate(values) if v is not None
            )
            fmt = "<c:formatCode>General</c:formatCode>" if tag == "numCache" else ""
            return f'<c:{tag}>{fmt}<c:ptCount val="{len(values)}"/>{points}</c:{tag}>'

        series = []
        for i, (name, categories, values) in enumerate(self.series):
            column = chr(ord("B") + i)
            series.append(
                f'<c:ser><c:idx val="{i}"/><c:order val="{i}"/>'
                f"<c:tx><c:strRef><c:f>Sheet1!${column}$1</c:f>{cache('strCache', [name])}</c:strRef></c:tx>"
                f"<c:cat><c:strRef><c:f>Sheet1!$A$2:$A${len(categories) + 1}</c:f>{cache('strCache', categories)}</c:strRef></c:cat>"
                f"<c:val><c:numRef><c:f>Sheet1!${column}$2:${column}${len(values) + 1}</c:f>{cache('numCache', values)}</c:numRef></c:val>"
                f"</c:ser>"
            )
        axes = '<c:axId val="1001"/><c:axId val="1002"/>' if self.chart_type != "pie" else ""
        plot = {
            "bar": '<c:barDir val="col"/><c:grouping val="clustered"/><c:varyColors val="0"/>',
            "line": '<c:grouping val="standard"/><c:varyColors val="0"/>',
            "pie": '<c:varyColors val="1"/>',
        }[self.chart_type]
        marker = '<c:marker val="1"/>' if self.chart_type == "line" else ""
        axis_xml = (
            '<c:catAx><c:axId val="1001"/><c:scaling><c:orientation val="minMax"/></c:scaling><c:delete val="0"/>'
            '<c:axPos val="b"/><c:crossAx val="1002"/></c:catAx>'
            '<c:valAx><c:axId val="1002"/><c:scaling><c:orientation val="minMax"/></c:scaling><c:delete val="0"/>'
            '<c:axPos val="l"/><c:crossAx val="1001"/></c:valAx>'
            if self.chart_type != "pie"
            else ""
        )
        title = (
            f'<c:title><c:tx><c:rich><a:bodyPr/><a:p><a:r><a:t>{escape(self.title)}</a:t></a:r></a:p></c:rich></c:tx>'
            f'<c:overlay val="0"/></c:title><c:autoTitleDeleted val="0"/>'
            if self.title is not None
            else '<c:autoTitleDeleted val="1"/>'
        )
        return (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            f'<c:chartSpace {nsdecls("c", "a", "r")}><c:chart>{title}<c:plotArea><c:layout/>'
            f'<c:{self.chart_type}Chart>{plot}{"".join(series)}{marker}{axes}</c:{self.chart_type}Chart>'
            f'{axis_xml}</c:plotArea><c:plotVisOnly val="1"/></c:chart></c:chartSpace>'
        )


def png(rgb: tuple[int, int, int]) -> bytes:
    """A 2x2 PNG of one color, written by hand so the script needs no imaging library."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    rows = b"".join(b"\x00" + bytes(rgb) * 2 for _ in range(2))
    header = struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b"")


def emf() -> bytes:
    """The smallest EMF: a header record and an end-of-file record, 1 mm square."""
    header = struct.pack(
        "<II4i4iIIIIHHIII4i",
        1, 88,  # EMR_HEADER, its size
        0, 0, 10, 10,  # bounds, device units
        0, 0, 100, 100,  # frame, 0.01 mm
        0x464D4520, 0x10000, 108, 2,  # " EMF", version, file size, records
        1, 0, 0, 0, 0,  # handles, reserved, no description, no palette
        1024, 768, 320, 240,  # reference device, pixels and millimeters
    )
    eof = struct.pack("<IIIII", 14, 20, 0, 16, 20)
    return header + eof


def headings() -> Fixture:
    f = Fixture("headings")
    f.heading("Quarterly Operations Report", 1)
    f.paragraph("This report summarizes operations for the third quarter.")
    f.heading("Scope", 2)
    f.paragraph("The scope covers three regional offices.")
    f.heading("Regional detail", 3)
    f.paragraph_with_break("North office: stable.", "South office: expanding.")
    f.heading("Notes on method", 4)
    f.paragraph("Figures marked * are estimates; see [appendix] for the _raw_ data | 2024 & later.")
    f.outline_heading("Outline-level heading", 2)
    f.paragraph("1. This paragraph starts like a numbered item but is plain text.")
    f.paragraph("# Not a heading either.")
    return f


def lists() -> Fixture:
    f = Fixture("lists")
    f.paragraph("Bulleted, nested three deep:")
    f.item("Fruit", 10, 0, False)
    f.item("Apples", 10, 1, False)
    f.item("Granny Smith", 10, 2, False)
    f.item("Pears", 10, 1, False)
    f.item("Vegetables", 10, 0, False)
    f.paragraph("Ordered, with a nested bullet list:")
    f.item("Prepare the site", 11, 0, True)
    f.item("Clear debris", 11, 1, False)
    f.item("Mark boundaries", 11, 1, False)
    f.item("Pour the foundation", 11, 0, True)
    f.item("Cure for seven days", 11, 0, True)
    f.paragraph("A second ordered list restarts its numbering:")
    f.item("Restart one", 12, 0, True)
    f.item("Restart two", 12, 0, True)
    f.item("Bullet under restart", 12, 1, False)
    f.item("Lettered detail", 12, 2, True)
    f.paragraph("Two ordered lists back to back, the second restarting, then a bullet list:")
    f.item("Back one", 13, 0, True)
    f.item("Back two", 13, 0, True)
    f.item("Again one", 14, 0, True, new_list=True)
    f.item("Again two", 14, 0, True)
    f.item("Bullet right after", 10, 0, False, new_list=True)
    f.paragraph("End of lists.")
    return f


def tables() -> Fixture:
    f = Fixture("tables")
    f.heading("Tables", 1)
    f.table(
        [
            [("Region", 1, 1), ("Q3", 1, 1), ("Q4", 1, 1)],
            [("North\n(provisional)", 1, 1), ("Pending audit", 2, 1), None],
            [("West", 1, 2), ("12", 1, 1), ("14", 1, 1)],
            [None, ("9", 1, 1), ("11 | revised", 1, 1)],
        ]
    )
    f.paragraph("Text after the table.")
    return f


def tracked_changes() -> Fixture:
    f = Fixture("tracked-changes")
    f.heading("Tracked changes", 1)
    f.tracked(
        [
            ("text", "The committee "),
            ("ins", "carefully "),
            ("text", "reviewed the proposal"),
            ("del", " and the obsolete appendix"),
            ("text", "."),
        ]
    )
    f.tracked(
        [
            ("text", "Budget figures follow. "),
            ("moveFrom", "This sentence moves down. "),
            ("text", "Staffing is unchanged."),
        ]
    )
    f.tracked([("text", "Closing remarks. "), ("moveTo", "This sentence moves down.")])
    f.tracked_joined([("text", "This sentence was split ")], [("text", "across two paragraphs.")])
    f.tracked_joined(
        [("del", "An entirely withdrawn paragraph.")],
        [("text", "The paragraph after the removed one.")],
    )
    f.tracked([("ins", "A newly inserted paragraph.")], mark="ins")
    f.tracked_item([("text", "Kept item"), ("del", " with withdrawn words")], 10, 0, False)
    f.tracked_item([("ins", "Inserted item")], 10, 0, False)
    f.tracked_before_table([("text", "A paragraph whose mark was deleted, right before a table.")])
    f.tracked_table(
        [
            [("text", "Cell "), ("del", "stale"), ("ins", "fresh"), ("text", " value")],
            [("text", "Plain cell")],
        ]
    )
    f.absent = ["obsolete appendix", "withdrawn", "stale"]
    f.once = ["This sentence moves down"]
    return f


def libreoffice_alt(truth: dict) -> None:
    """LibreOffice writes a picture's title and description as one description, the title's line
    first, and no title."""

    def walk(node) -> None:
        if isinstance(node, dict):
            if node.get("kind") == "media" and "title" in node:
                title = node.pop("title")
                node["alt"] = f"{title}\n{node['alt']}" if "alt" in node else title
                # Keep the key order of the python-docx truth: alt before section.
                node["section"] = node.pop("section")
            for value in node.values():
                walk(value)
        elif isinstance(node, list):
            for value in node:
                walk(value)

    walk(truth)


def resection(elements: list[dict], section: int) -> list[dict]:
    """A copy of a story's ground truth read for another section: its media placed there."""
    copy = json.loads(json.dumps(elements))
    for element in copy:
        if element["kind"] == "media":
            element["section"] = section
    return copy


def story_truth(**kinds: list[dict]) -> list[dict]:
    return [{"kind": kind, "elements": elements} for kind, elements in kinds.items()]


def sections() -> Fixture:
    """Three sections, with even-page headers on for the document. The first has its own default,
    first-page, and even-page headers and footers and a title page; the second its own default
    and first-page headers, no title page (so its first-page header doesn't count), and inherits
    the rest; the third has a title page and inherits everything, the first-page header the
    second section's."""
    f = Fixture("sections")
    f.doc.settings.odd_and_even_pages_header_footer = True
    logo = Picture("Header logo", "Logo", png((200, 30, 30)))

    s0 = f.doc.sections[0]
    s0.different_first_page_header_footer = True
    with f.story(s0.header, 0) as h0_default:
        f.paragraph("Default header one")
    # The tracked change is in a header no other section inherits: LibreOffice writes an
    # inherited header's copies with its deletions as plain text.
    with f.story(s0.first_page_header, 0) as h0_first:
        f.tracked([("text", "First header "), ("ins", "one"), ("del", " withdrawn")])
    with f.story(s0.even_page_header, 0) as h0_even:
        f.paragraph("Even header one")
    with f.story(s0.footer, 0) as f0_default:
        f.table([[("Footer cell A", 1, 1), ("Footer cell B", 1, 1)]])
    with f.story(s0.first_page_footer, 0) as f0_first:
        f.paragraph("First footer one")
    with f.story(s0.even_page_footer, 0) as f0_even:
        f.item("Even footer item", 10, 0, False)
    f.heading("Section one", 1)
    f.paragraph("Body of section one.")

    s1 = f.new_section()
    s1.different_first_page_header_footer = False
    # python-docx writes a linked header's content into the section it inherits from.
    s1.header.is_linked_to_previous = False
    s1.first_page_header.is_linked_to_previous = False
    with f.story(s1.header, 1) as h1_default:
        f.paragraph("Default header two")
        f.tracked([("picture", logo)])
    with f.story(s1.first_page_header, 1) as h1_first:
        f.paragraph("First header two")
    f.heading("Section two", 1)
    f.paragraph("Body of section two.")

    s2 = f.new_section()
    s2.different_first_page_header_footer = True
    f.heading("Section three", 1)
    f.paragraph("Body of section three.")

    f.sections = [
        {
            "headers": story_truth(default=h0_default, first=h0_first, even=h0_even),
            "footers": story_truth(default=f0_default, first=f0_first, even=f0_even),
        },
        {
            "headers": story_truth(default=h1_default, even=resection(h0_even, 1)),
            "footers": story_truth(default=resection(f0_default, 1), even=resection(f0_even, 1)),
        },
        {
            "headers": story_truth(default=resection(h1_default, 2), first=resection(h1_first, 2), even=resection(h0_even, 2)),
            "footers": story_truth(default=resection(f0_default, 2), first=resection(f0_first, 2), even=resection(f0_even, 2)),
        },
    ]
    f.absent = ["withdrawn"]

    def libreoffice(truth: dict) -> None:
        # LibreOffice gives the third section the default header and footer as its first-page
        # ones: it drops the first-page variants the second section, with no title page, passes on.
        libreoffice_alt(truth)
        for story in ("headers", "footers"):
            kinds = {s["kind"]: s for s in truth["sections"][2][story]}
            kinds["first"]["elements"] = kinds["default"]["elements"]

    f.libreoffice = libreoffice
    return f


def headers_even_off() -> Fixture:
    """First-page and even-page headers written, but no title page and no even-page headers in
    the settings, so only the default header and footer count."""
    f = Fixture("headers-even-off")
    s0 = f.doc.sections[0]
    with f.story(s0.header, 0) as header:
        f.paragraph("Plain header")
    s0.first_page_header.add_paragraph("Unused first header")
    s0.even_page_header.add_paragraph("Unused even header")
    with f.story(s0.footer, 0) as footer:
        f.paragraph("Plain footer")
    f.paragraph("Body text without even-page or first-page variants.")
    f.sections = [{"headers": story_truth(default=header), "footers": story_truth(default=footer)}]
    f.absent = ["Unused first header", "Unused even header"]
    return f


def notes() -> Fixture:
    """Footnotes and endnotes, referenced out of their id order, one note of two paragraphs, a
    reference in a list item, and a footnote whose only reference was deleted."""
    f = Fixture("notes")
    f.note("footnote", "basis", ["Figures are unaudited."])
    f.note("footnote", "method", ["The method has two parts.", "Its second part is a second paragraph."])
    f.note("footnote", "withdrawn", ["A withdrawn footnote."])
    f.note("endnote", "source", ["Source: the regional ledgers."])
    f.note("endnote", "item", ["An endnote from a list item."])
    f.heading("Notes", 1)
    f.tracked(
        [
            ("text", "Revenue rose sharply"),
            ("footnote", "method"),
            ("text", " in the third quarter."),
            ("endnote", "source"),
        ]
    )
    f.tracked([("text", "Costs held steady"), ("footnote", "basis"), ("delNote", ("footnote", "withdrawn")), ("text", ".")])
    f.tracked_item([("text", "An item with an endnote"), ("endnote", "item")], 10, 0, False)
    f.absent = ["A withdrawn footnote"]

    def libreoffice(truth: dict) -> None:
        # LibreOffice pairs footnotes with their references by their order in the notes part, not
        # by id, so in its re-save the first reference shows the first-defined note.
        method, basis = truth["footnotes"]
        truth["footnotes"] = [basis, method]

    f.libreoffice = libreoffice
    return f


def comments() -> Fixture:
    """Comments with commentsExtended: a thread of two replies, a resolved comment spanning two
    paragraphs with two paragraphs of text, and a comment over tracked text."""
    f = Fixture("comments")
    f.comment(
        "budget",
        "Ana Reviewer",
        ["Please check these figures."],
        "budget figures",
        replies=[("Ben Author", ["Checked against the ledger."]), ("Ana Reviewer", ["Thanks."])],
    )
    f.comment(
        "span",
        "Ben Author",
        ["This comment spans two paragraphs.", "Its text has two paragraphs as well."],
        "spans two\nparagraphs",
        resolved=True,
    )
    f.comment("tracked", "Cleo Editor", ["An open comment over tracked text."], "timeline")
    f.heading("Comments", 1)
    f.tracked(
        [
            ("text", "The "),
            ("commentStart", "budget"),
            ("text", "budget figures"),
            ("commentEnd", "budget"),
            ("text", " need review."),
        ]
    )
    f.tracked([("text", "A comment that "), ("commentStart", "span"), ("text", "spans two")])
    f.tracked([("text", "paragraphs"), ("commentEnd", "span"), ("text", " of text.")])
    f.tracked(
        [
            ("text", "The "),
            ("commentStart", "tracked"),
            ("text", "time"),
            ("del", " "),
            ("ins", "line"),
            ("commentEnd", "tracked"),
            ("text", " is new."),
        ]
    )
    return f


def comments_plain() -> Fixture:
    """A comment written by python-docx, with no commentsExtended part: no replies, unresolved."""
    f = Fixture("comments-plain")
    p = f.doc.add_paragraph("Before ")
    anchor = p.add_run("the anchor")
    p.add_run(" and after.")
    f.doc.add_comment(anchor, text="A comment without commentsExtended.", author="Dee Writer", initials="DW")
    text = "Before the anchor and after."
    f.elements.append({"kind": "paragraph", "text": text, "marks": [{"at": len("Before the anchor"), "kind": "comment", "n": 1}]})
    f.plain_comments.append(
        {
            "author": "Dee Writer",
            "anchor": "the anchor",
            "resolved": False,
            "elements": [{"kind": "paragraph", "text": "A comment without commentsExtended."}],
            "replies": [],
        }
    )
    return f


def media() -> Fixture:
    """Pictures: one on its own line, one between text, the first's bytes again through the same
    part and through a part of its own (one file either way), a metafile, a linked picture, a
    deleted picture, one in a table cell, and one in a second section."""
    f = Fixture("media")
    red, green, blue = png((220, 40, 40)), png((40, 160, 60)), png((40, 60, 200))
    f.heading("Media", 1)
    f.paragraph("A picture on its own line follows.")
    f.tracked([("picture", Picture("A red square", "Red", red))])
    f.tracked([("text", "Text before "), ("picture", Picture("A blue square", "", blue)), ("text", "and after.")])
    f.tracked([("picture", Picture("The red square again", "", red))])
    f.tracked([("picture", Picture("The red square in a part of its own", "", red, part="/word/media/copy-of-red.png"))])
    f.tracked([("picture", Picture("A metafile", "", emf(), "emf", CT.X_EMF, part="/word/media/vector.emf"))])
    f.tracked([("picture", Picture("A linked picture", "Linked", link="https://example.invalid/linked.png"))])
    f.tracked([("text", "A deleted picture follows."), ("delPicture", Picture("Withdrawn picture", "", green))])
    f.tracked_table(
        [
            [("text", "Cell with a picture"), ("picture", Picture("A green square", "", green))],
            [("text", "Plain cell")],
        ]
    )
    f.new_section()
    f.tracked([("picture", Picture("Blue in section two", "", blue))])
    f.absent = ["Withdrawn picture"]
    f.libreoffice = libreoffice_alt
    return f


def charts() -> Fixture:
    """Charts as cached data: a clustered bar chart of two series with a missing value, an
    untitled pie chart, and a line chart."""
    f = Fixture("charts")
    f.heading("Charts", 1)
    f.paragraph("Revenue by region:")
    f.tracked(
        [
            (
                "chart",
                Chart(
                    "bar",
                    "Revenue by region",
                    [("North", ["Q1", "Q2", "Q3"], [12.5, 14, 15.25]), ("South", ["Q1", "Q2", "Q3"], [9, None, 11])],
                ),
            )
        ]
    )
    f.paragraph("Share of costs:")
    f.tracked([("chart", Chart("pie", None, [("Share", ["Staff", "Rent", "Other"], [60, 25, 15])]))])
    f.tracked([("chart", Chart("line", "Trend", [("Index", ["Jan", "Feb", "Mar", "Apr"], [1.5, 2, 2.75, 3.25])]))])
    return f


FIXTURES = [headings, lists, tables, tracked_changes, sections, headers_even_off, notes, comments, comments_plain, media, charts]


# xlsx: workbooks written by openpyxl and re-saved by LibreOffice.

CT_PNG = "image/png"


class RawImage(Image):
    """A picture openpyxl stores as given. Its own Image reads the picture through Pillow, which
    this script has no need of."""

    def __init__(self, data: bytes):
        self.ref = data
        self.width = self.height = 48
        self.format = "png"

    def _data(self) -> bytes:
        return self.ref


VARIANTS = ("openpyxl", "libreoffice")


class Workbook:
    """An openpyxl workbook and its ground truth, built side by side. A cell's truth is kept per
    writer, as each leaves it: LibreOffice computes the formulas openpyxl leaves uncached, writes
    constant booleans and errors as formulas, and expands shared formulas into their cells'
    own."""

    writer, extension, filter = "openpyxl", "xlsx", "xlsx:Calc MS Excel 2007 XML"

    def __init__(self, name: str, date1904: bool = False):
        self.name = name
        self.book = openpyxl.Workbook()
        self.book.remove(self.book.active)
        if date1904:
            self.book.epoch = CALENDAR_MAC_1904
        self.sheets: list[Sheet] = []
        # Raw XML for what openpyxl can't write: (part, the XML openpyxl writes, the XML instead).
        self.patches: list[tuple[str, str, str]] = []

    def sheet(self, title: str, state: str = "visible") -> "Sheet":
        ws = self.book.create_sheet(title)
        ws.sheet_state = state
        sheet = Sheet(self, ws, len(self.sheets))
        self.sheets.append(sheet)
        return sheet

    def save(self, path: pathlib.Path) -> None:
        # openpyxl numbers a drawing part for each sheet that has charts or pictures, in sheet
        # order, and in each names its charts and then its pictures by their count from 1.
        drawing = 0
        for sheet in self.sheets:
            if not sheet.charts and not sheet.pictures:
                continue
            drawing += 1
            for i, (_, alt) in enumerate(sheet.pictures, len(sheet.charts) + 1):
                self.patches.append(
                    (f"xl/drawings/drawing{drawing}.xml", f'name="Image {i}" descr="Picture"', f'name="Image {i}" descr="{escape(alt)}"')
                )
        self.book.save(path)
        with zipfile.ZipFile(path) as source:
            entries = [(info, source.read(info.filename)) for info in source.infolist()]
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
            for info, data in entries:
                for part, old, new in self.patches:
                    if part == info.filename:
                        text = data.decode()
                        if text.count(old) != 1:
                            raise RuntimeError(f"{self.name}: {part} holds {old!r} {text.count(old)} times, not once")
                        data = text.replace(old, new).encode()
                target.writestr(info, data)

    def _truth(self, variant: str) -> dict:
        return {
            "fixture": self.name,
            "format": "xlsx",
            "elements": [sheet.truth(variant) for sheet in self.sheets],
            "absent": [],
            "once": [],
        }

    def truth(self) -> dict:
        return self._truth("openpyxl")

    def libreoffice_truth(self) -> dict | None:
        adjusted = self._truth("libreoffice")
        return None if adjusted == self.truth() else adjusted


class Sheet:
    """A worksheet and its truth: its cells by reference, each writer's truth of each (None where
    the writer leaves the cell out), then its charts and its pictures."""

    def __init__(self, book: Workbook, ws, index: int):
        self.book = book
        self.ws = ws
        self.index = index
        self.cells: dict[str, dict[str, dict | None]] = {}
        self.charts: list[dict[str, dict]] = []
        self.pictures: list[tuple[dict, str]] = []

    @property
    def part(self) -> str:
        return f"xl/worksheets/sheet{self.index + 1}.xml"

    def _cell(self, ref: str, openpyxl_truth: dict | None, libreoffice_truth: dict | None) -> None:
        self.cells[ref] = {"openpyxl": openpyxl_truth, "libreoffice": libreoffice_truth}

    def value(self, ref: str, value, kind: str, expected=None, number_format: str | None = None, libreoffice: dict | None = None):
        """A constant cell: its value, its truth's type and value (the value itself unless
        expected says otherwise), and the properties LibreOffice's re-save adds to its truth."""
        self.ws[ref] = value
        if number_format is not None:
            self.ws[ref].number_format = number_format
        truth = {"ref": ref, "type": kind, "value": value if expected is None else expected}
        self._cell(ref, truth, {**truth, **(libreoffice or {})})

    def styled(self, ref: str, number_format: str) -> None:
        """A cell with a style and no value, which both writers keep and the truth leaves out."""
        self.ws[ref].number_format = number_format

    def formula(self, ref: str, text: str, kind: str, cached, number_format: str | None = None) -> None:
        """A formula: openpyxl writes it with an empty cached value, LibreOffice with the value it
        computes (kind and cached)."""
        self.ws[ref] = "=" + text
        if number_format is not None:
            self.ws[ref].number_format = number_format
        self._cell(ref, {"ref": ref, "formula": text}, {"ref": ref, "type": kind, "value": cached, "formula": text})

    def shared_formula(self, cells: list[tuple[str, str, str, object]], index: int = 0) -> None:
        """A shared formula over cells of (reference, the cell's own formula, the type and value
        LibreOffice computes), the first the anchor holding the text. openpyxl writes each
        cell's own formula, patched here into the shared one; LibreOffice writes them back as
        their own formulas."""
        anchor = cells[0][0]
        span = f"{anchor}:{cells[-1][0]}"
        for n, (ref, text, kind, cached) in enumerate(cells):
            self.ws[ref] = "=" + text
            shared = f'<f t="shared" ref="{span}" si="{index}">{escape(text)}</f>' if n == 0 else f'<f t="shared" si="{index}"/>'
            self.book.patches.append((self.part, f'<c r="{ref}"><f>{escape(text)}</f><v></v></c>', f'<c r="{ref}">{shared}<v></v></c>'))
            own = {"ref": ref, "formula": text} if n == 0 else {"ref": ref, "sharedFrom": anchor}
            self._cell(ref, own, {"ref": ref, "type": kind, "value": cached, "formula": text})

    def array_formula(self, span: str, text: str, cached: list[tuple[str, str, object]]) -> None:
        """An array formula over span, its text on the anchor: openpyxl writes the anchor alone,
        LibreOffice each cell's computed (reference, type, value), the anchor's with the text."""
        anchor = span.split(":")[0]
        self.ws[anchor] = ArrayFormula(span, "=" + text)
        for n, (ref, kind, value) in enumerate(cached):
            computed = {"ref": ref, "type": kind, "value": value}
            if n == 0:
                self._cell(ref, {"ref": ref, "formula": text}, {**computed, "formula": text})
            else:
                self._cell(ref, None, computed)

    def raw(self, ref: str, xml: str, truth: dict, libreoffice: dict) -> None:
        """A cell openpyxl can't write, patched in as raw XML over a placeholder string."""
        self.ws[ref] = f"placeholder {ref}"
        self.book.patches.append((self.part, f'<c r="{ref}" t="inlineStr"><is><t>placeholder {ref}</t></is></c>', xml))
        self._cell(ref, {"ref": ref, **truth}, {"ref": ref, **libreoffice})

    def merge(self, span: str, value, kind: str) -> None:
        """Merged cells, the value in the anchor; both writers keep the merge."""
        self.ws.merge_cells(span)
        first, last = (coordinate_to_tuple(c) for c in span.split(":"))
        anchor = span.split(":")[0]
        self.ws[anchor] = value
        truth = {"ref": anchor, "type": kind, "value": value, "colSpan": last[1] - first[1] + 1, "rowSpan": last[0] - first[0] + 1}
        self._cell(anchor, truth, dict(truth))

    def bar_chart(self, anchor: str, title: str, categories: str, values: str, series: tuple[str, list[str], list[float]]) -> None:
        """A bar chart of one series from the sheet's cells: categories and values are ranges of
        one column, values with its name's cell first. openpyxl writes the references with no
        cached data, so its chart has no series name, categories, or values; LibreOffice caches
        them (series)."""
        chart = BarChart()
        chart.title = title
        chart.add_data(Reference(self.ws, range_string=f"'{self.ws.title}'!{values}"), titles_from_data=True)
        chart.set_categories(Reference(self.ws, range_string=f"'{self.ws.title}'!{categories}"))
        self.ws.add_chart(chart, anchor)
        name, cats, vals = series
        base = {"kind": "chart", "chartType": "bar", "title": title}
        self.charts.append(
            {
                "openpyxl": {**base, "series": [{"name": None, "categories": [], "values": []}]},
                "libreoffice": {**base, "series": [{"name": name, "categories": cats, "values": vals}]},
            }
        )

    def picture(self, anchor: str, data: bytes, alt: str) -> None:
        """A PNG picture, each its own part in openpyxl's file; LibreOffice keeps one part for
        the same bytes. Its alternative text is patched in over openpyxl's "Picture"."""
        image = RawImage(data)
        image.anchor = anchor
        self.ws.add_image(image)
        truth = {
            "kind": "media",
            "name": hashlib.sha256(data).hexdigest() + ".png",
            "contentType": CT_PNG,
            "alt": alt,
            "sheet": self.index,
        }
        self.pictures.append((truth, alt))

    def truth(self, variant: str) -> dict:
        def position(ref: str) -> tuple[int, int]:
            return coordinate_to_tuple(ref)

        rows: dict[int, list[dict]] = {}
        for ref in sorted(self.cells, key=position):
            cell = self.cells[ref][variant]
            if cell is not None:
                rows.setdefault(position(ref)[0], []).append(cell)
        elements = []
        if rows:
            elements.append({"kind": "table", "rows": [{"row": r, "cells": cells} for r, cells in rows.items()]})
        elements += [chart[variant] for chart in self.charts]
        elements += [picture for picture, _ in self.pictures]
        return {
            "kind": "sheet",
            "name": self.ws.title,
            "hidden": self.ws.sheet_state != "visible",
            "elements": elements,
        }


def cells() -> Workbook:
    """Constants of each type, through each writer's strings (openpyxl writes them inline,
    LibreOffice shared), rich text, date, datetime, and time numbers by built-in and custom
    formats, numbers in formats that only look like dates, ISO dates stored as such, a styled
    cell with no value, and a sparse corner."""
    f = Workbook("cells")
    s = f.sheet("Cells")
    labels = [
        "Text", "Rich text", "Numbers", "Booleans", "Errors", "Percent", "Thousands", "Date", "Datetime",
        "Time", "Built-in date", "Built-in datetime", "Built-in time", "Custom date", "Custom time",
        "Quoted letters", "Bracketed letters", "Escaped letters", "Early dates", "ISO date",
    ]
    for row, label in enumerate(labels, 1):
        s.value(f"A{row}", label, "text")
    s.value("B1", "plain text", "text")
    s.value("B2", CellRichText(["plain ", TextBlock(InlineFont(b=True), "bold")]), "text", expected="plain bold")
    s.value("B3", 12.5, "number")
    s.value("C3", -3, "number")
    s.value("D3", 1234.5678, "number")
    s.value("B4", True, "boolean", libreoffice={"formula": "TRUE()"})
    s.value("C4", False, "boolean", libreoffice={"formula": "FALSE()"})
    s.value("B5", "#N/A", "error", libreoffice={"formula": "#N/A"})
    s.value("C5", "#DIV/0!", "error", libreoffice={"formula": "#DIV/0!"})
    s.value("B6", 0.125, "number", number_format="0.00%")
    s.value("B7", 1234567, "number", number_format="#,##0")
    s.value("B8", datetime.date(2026, 1, 15), "date", expected="2026-01-15")
    s.value("B9", datetime.datetime(2026, 1, 15, 10, 30), "date", expected="2026-01-15T10:30:00")
    s.value("B10", datetime.time(14, 15, 30), "date", expected="14:15:30")
    s.value("B11", datetime.date(2025, 12, 31), "date", expected="2025-12-31", number_format="mm-dd-yy")
    s.value("B12", datetime.datetime(2025, 6, 1, 8, 5), "date", expected="2025-06-01T08:05:00", number_format="m/d/yy h:mm")
    s.value("B13", datetime.time(9, 45), "date", expected="09:45:00", number_format="h:mm")
    s.value("B14", datetime.date(2026, 3, 1), "date", expected="2026-03-01", number_format='dddd", "mmmm d", "yyyy')
    s.value("B15", datetime.time(18, 0), "date", expected="18:00:00", number_format="h:mm AM/PM")
    s.value("B16", 42, "number", number_format='0 "days"')
    s.value("B17", 7, "number", number_format="[Red]0;[Blue]-0")
    s.value("B18", 5, "number", number_format="0\\h")
    # Either side of the 1900 system's phantom 29 February.
    s.value("B19", datetime.date(1900, 2, 28), "date", expected="1900-02-28")
    s.value("C19", datetime.date(1900, 3, 1), "date", expected="1900-03-01")
    s.raw(
        "B20",
        '<c r="B20" t="d"><v>2026-01-15T10:30:00</v></c>',
        {"type": "date", "value": "2026-01-15T10:30:00"},
        # LibreOffice stores it as the serial number, in the General format.
        {"type": "number", "value": 46037.4375},
    )
    s.styled("C22", "0.00")
    s.value("E23", "far corner", "text")
    return f


def formulas() -> Workbook:
    """Formulas of each result type, uncached in openpyxl's file and cached in LibreOffice's, a
    shared formula (raw XML; LibreOffice expands it), an array formula, and a formula whose
    value is a date by its format."""
    f = Workbook("formulas")
    s = f.sheet("Formulas")
    for row in (1, 2, 3):
        s.value(f"A{row}", row, "number")
        s.value(f"B{row}", row * 10, "number")
    s.shared_formula([(f"C{r}", f"A{r}*B{r}", "number", r * r * 10) for r in (1, 2, 3)])
    s.array_formula("D1:D3", "A1:A3*2", [("D1", "number", 2), ("D2", "number", 4), ("D3", "number", 6)])
    s.formula("A5", "SUM(A1:A3)", "number", 6)
    s.formula("A6", 'A1&" unit"', "text", "1 unit")
    s.formula("A7", "A1/0", "error", "#DIV/0!")
    s.formula("A8", "A1>B1", "boolean", False)
    s.formula("A9", "DATE(2026,1,15)", "date", "2026-01-15", number_format="yyyy-mm-dd")
    return f


def sheets() -> Workbook:
    """Sheets in an order that isn't their names', hidden and very hidden ones, an empty one,
    and merged cells across columns, rows, and both."""
    f = Workbook("sheets")
    f.sheet("Zeta").value("A1", "the first sheet", "text")
    f.sheet("Hidden", "hidden").value("A1", "a hidden sheet", "text")
    merged = f.sheet("Alpha")
    merged.merge("A1:C1", "Title", "text")
    merged.merge("A2:A3", "Rows", "text")
    merged.value("B2", 1, "number")
    merged.value("B3", 2, "number")
    merged.value("C2", "beside", "text")
    merged.merge("B5:C6", "Block", "text")
    f.sheet("Empty")
    f.sheet("Very hidden", "veryHidden").value("A1", "a very hidden sheet", "text")
    return f


def drawings() -> Workbook:
    """A sheet's chart and pictures: a bar chart over the sheet's cells, the same picture twice
    (two parts in openpyxl's file, one file either way), and another sheet's picture."""
    f = Workbook("drawings")
    red, blue = png((220, 40, 40)), png((40, 60, 200))
    s = f.sheet("Sales")
    s.value("A1", "Quarter", "text")
    s.value("B1", "Revenue", "text")
    for row, (quarter, revenue) in enumerate([("Q1", 12.5), ("Q2", 14), ("Q3", 15.25)], 2):
        s.value(f"A{row}", quarter, "text")
        s.value(f"B{row}", revenue, "number")
    s.bar_chart("D2", "Revenue by quarter", "A2:A4", "B1:B4", ("Revenue", ["Q1", "Q2", "Q3"], [12.5, 14, 15.25]))
    s.picture("D18", red, "A red square")
    s.picture("H18", red, "The red square again")
    other = f.sheet("Other")
    other.value("A1", "A picture follows.", "text")
    other.picture("B2", blue, "A blue square")
    return f


def dates_1904() -> Workbook:
    """Dates in the 1904 date system, whose serial 0 is 1 January 1904."""
    f = Workbook("dates-1904", date1904=True)
    s = f.sheet("Dates")
    s.value("A1", datetime.date(2026, 1, 15), "date", expected="2026-01-15")
    s.value("A2", datetime.datetime(1904, 1, 1, 12, 0), "date", expected="1904-01-01T12:00:00")
    s.value("A3", datetime.date(1904, 1, 2), "date", expected="1904-01-02")
    s.value("A4", 1.5, "number")
    return f


XLSX_FIXTURES = [cells, formulas, sheets, drawings, dates_1904]


# pptx: presentations written by python-pptx and re-saved by LibreOffice.

TITLE_SLIDE, TITLE_AND_CONTENT, TITLE_ONLY, BLANK = 0, 1, 5, 6
A = nsdecls("a")
# The prompt texts of python-pptx's default layouts and master, which no slide holds, and the
# slide-number field LibreOffice writes in each notes page's slide-number placeholder.
ABSENT_ON_SLIDES = ["Click to edit Master title style", "Click to edit Master text styles", "Click to edit Master subtitle style", "<number>"]


def bullet_xml(kind: str) -> str:
    """A paragraph's own bullet: "number" an arabic auto-number, "char" a character bullet,
    "none" no bullet; "inherit" writes none, so the placeholder's layout or master decides."""
    return {
        "number": f'<a:buFont {A} typeface="+mj-lt"/><a:buAutoNum {A} type="arabicPeriod"/>',
        "char": f'<a:buFont {A} typeface="Arial"/><a:buChar {A} char="•"/>',
        "none": f"<a:buNone {A}/>",
        "inherit": "",
    }[kind]


def list_truth(items: list[tuple[str, int, str]], listed) -> list[dict]:
    """The truth of a text body's paragraphs of (text, level, bullet kind): a paragraph whose
    bullet (its own, or inherited where listed says it is listed) is none is a paragraph; the
    others are list items, nested by level, a new list starting where the one before is of the
    other kind at the same level or a paragraph came between. Levels never skip, so an item's
    nesting is its level."""
    elements, open_lists = [], []  # open_lists: each open level's ordered flag
    for text, level, kind in items:
        ordered = {"number": True, "char": False, "none": None, "inherit": listed(level)}[kind]
        if ordered is None:
            open_lists = []
            elements.append({"kind": "paragraph", "text": text})
            continue
        del open_lists[level + 1 :]
        start = len(open_lists) <= level or open_lists[level] != ordered
        open_lists[level:] = [ordered]
        elements.append({"kind": "listItem", "level": level, "ordered": ordered, "start": start, "text": text})
    return elements


def master_bullets(level: int) -> bool | None:
    """The default master's body style: a character bullet at every level, so unordered."""
    return False


def layout_none(level: int) -> bool | None:
    """A layout placeholder whose list style turns bullets off (the title slide's subtitle)."""
    return None


class Deck:
    """A python-pptx presentation and its ground truth, built side by side."""

    writer, extension, filter = "python-pptx", "pptx", "pptx:Impress MS PowerPoint 2007 XML"

    def __init__(self, name: str):
        self.name = name
        self.prs = pptx.Presentation()
        self.slides: list["DeckSlide"] = []
        self.order: list[int] | None = None
        self.absent: list[str] = list(ABSENT_ON_SLIDES)
        self.once: list[str] = []
        # How the LibreOffice re-save's ground truth differs, when it does.
        self.libreoffice = None

    def slide(self, layout: int, hidden: bool = False) -> "DeckSlide":
        s = DeckSlide(self, self.prs.slides.add_slide(self.prs.slide_layouts[layout]), hidden)
        if hidden:
            s.slide._element.set("show", "0")
        self.slides.append(s)
        return s

    def reorder(self, order: list[int]) -> None:
        """Presents the slides in the given order of their creation, so the slide parts' names
        (slide1.xml, ...) aren't the presentation's order."""
        ids = self.prs.slides._sldIdLst
        entries = list(ids)
        for entry in entries:
            ids.remove(entry)
        for i in order:
            ids.append(entries[i])
        self.order = order

    def save(self, path: pathlib.Path) -> None:
        self.prs.save(path)

    def presented(self) -> list["DeckSlide"]:
        return [self.slides[i] for i in self.order] if self.order is not None else self.slides

    def truth(self) -> dict:
        return {
            "fixture": self.name,
            "format": "pptx",
            "elements": [s.truth(n) for n, s in enumerate(self.presented())],
            "absent": self.absent,
            "once": self.once,
        }

    def libreoffice_truth(self) -> dict | None:
        if self.libreoffice is None:
            return None
        adjusted = json.loads(json.dumps(self.truth()))
        self.libreoffice(adjusted)
        return adjusted


class DeckSlide:
    """A slide and its truth: its elements in shape-tree order, and its notes."""

    def __init__(self, deck: Deck, slide, hidden: bool):
        self.deck = deck
        self.slide = slide
        self.hidden = hidden
        self.elements: list = []  # truth, or a callable of the slide's index for media
        self.notes: list[dict] = []

    def title(self, text: str) -> None:
        self.slide.shapes.title.text = text
        self.elements.append({"kind": "heading", "level": 1, "text": text})

    def placeholder_text(self, idx: int, items: list[tuple[str, int, str]], listed) -> None:
        """The placeholder idx's paragraphs of (text, level, bullet kind)."""
        self._fill(self.slide.placeholders[idx].text_frame, items)
        self.elements += list_truth(items, listed)

    def textbox(self, items: list[tuple[str, int, str]], top: int = 0, shapes=None) -> None:
        """A text box, which no placeholder styles: its paragraphs are bulleted only by their own
        bullets."""
        box = (shapes if shapes is not None else self.slide.shapes).add_textbox(Inches(1), Inches(1) + top, Inches(4), Inches(1))
        self._fill(box.text_frame, items)
        self.elements += list_truth(items, lambda level: None)

    def group(self, build) -> None:
        """A group shape whose shapes build adds through the group's shape collection."""
        group = self.slide.shapes.add_group_shape()
        build(group.shapes)

    def notes_text(self, paragraphs: list[str]) -> None:
        frame = self.slide.notes_slide.notes_text_frame
        frame.text = paragraphs[0]
        for text in paragraphs[1:]:
            frame.add_paragraph().text = text
        self.notes = [{"kind": "paragraph", "text": p} for p in paragraphs]

    def table(self, rows: list[list[str | None]], merges: list[tuple[int, int, int, int]]) -> None:
        """A table of cells' text (None for a cell a merge covers, a newline between a cell's
        paragraphs) and merges of (row, column, rows, columns), python-pptx writing gridSpan and
        rowSpan on the origin and hMerge and vMerge on the cells it covers."""
        shape = self.slide.shapes.add_table(len(rows), len(rows[0]), Inches(1), Inches(2), Inches(6), Inches(2))
        table = shape.table
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                if text is not None:
                    frame = table.cell(r, c).text_frame
                    first, *rest = text.split("\n")
                    frame.text = first
                    for line in rest:
                        frame.add_paragraph().text = line
        spans = {}
        for r, c, height, width in merges:
            table.cell(r, c).merge(table.cell(r + height - 1, c + width - 1))
            spans[(r, c)] = (height, width)
        self.elements.append(
            {
                "kind": "table",
                "rows": [
                    [
                        {"text": text, "colSpan": spans.get((r, c), (1, 1))[1], "rowSpan": spans.get((r, c), (1, 1))[0]}
                        for c, text in enumerate(row)
                        if text is not None
                    ]
                    for r, row in enumerate(rows)
                ],
            }
        )

    def chart(self, title: str | None, categories: list[str], series: list[tuple[str, list[float | None]]]) -> None:
        """A clustered column chart, python-pptx writing its embedded workbook and its caches."""
        data = CategoryChartData()
        data.categories = categories
        for name, values in series:
            data.add_series(name, values)
        chart = self.slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), data).chart
        chart.has_title = title is not None
        if title is not None:
            chart.chart_title.text_frame.text = title
        element = {"kind": "chart", "chartType": "bar"}
        if title is not None:
            element["title"] = title
        element["series"] = [{"name": n, "categories": categories, "values": v} for n, v in series]
        self.elements.append(element)

    def picture(self, data: bytes, descr: str, title: str = "", shapes=None) -> None:
        """A PNG picture; python-pptx keeps one part for the same bytes, however often shown."""
        pic = (shapes if shapes is not None else self.slide.shapes).add_picture(io.BytesIO(data), Inches(1), Inches(4), Inches(1), Inches(1))
        properties = pic._element.nvPicPr.cNvPr
        properties.set("descr", descr)
        if title:
            properties.set("title", title)

        def truth(index: int) -> dict:
            element = {"kind": "media", "name": hashlib.sha256(data).hexdigest() + ".png", "contentType": CT_PNG, "alt": descr}
            if title:
                element["title"] = title
            element["slide"] = index
            return element

        self.elements.append(truth)

    @staticmethod
    def _fill(frame, items: list[tuple[str, int, str]]) -> None:
        for n, (text, level, kind) in enumerate(items):
            p = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
            p.text = text
            p.level = level
            if kind != "inherit":
                p_pr = p._p.get_or_add_pPr()
                if kind == "none":
                    p_pr.set("marL", "0")
                    p_pr.set("indent", "0")
                for element in parse_xml(f"<a:x {A}>{bullet_xml(kind)}</a:x>"):
                    p_pr.append(element)

    def truth(self, index: int) -> dict:
        element = {
            "kind": "slide",
            "number": index + 1,
            "hidden": self.hidden,
            "elements": [e(index) if callable(e) else e for e in self.elements],
        }
        if self.notes:
            element["notes"] = self.notes
        return element


def libreoffice_no_picture_title(truth: dict) -> None:
    """LibreOffice keeps a slide picture's description and drops its title."""

    def walk(node) -> None:
        if isinstance(node, dict):
            if node.get("kind") == "media":
                node.pop("title", None)
            for value in node.values():
                walk(value)
        elif isinstance(node, list):
            for value in node:
                walk(value)

    walk(truth)


def slides() -> Deck:
    """Slides presented in an order that isn't their parts', a hidden one, notes on some, a title
    slide's subtitle (its layout turns the master's bullets off), a text box, and a content
    placeholder left empty, whose layout prompt must not appear."""
    f = Deck("slides")
    welcome = f.slide(TITLE_SLIDE)
    welcome.title("Quarterly Review")
    welcome.placeholder_text(1, [("Operations, third quarter", 0, "inherit"), ("Prepared for the board", 0, "inherit")], layout_none)
    welcome.notes_text(["Welcome everyone.", "Introduce the agenda."])
    agenda = f.slide(TITLE_AND_CONTENT)
    agenda.title("Agenda")
    agenda.placeholder_text(1, [("Results", 0, "inherit"), ("Outlook", 0, "inherit")], master_bullets)
    internal = f.slide(TITLE_AND_CONTENT, hidden=True)
    internal.title("Internal figures")
    internal.placeholder_text(1, [("Not for the board", 0, "inherit")], master_bullets)
    internal.notes_text(["Skip this slide when presenting."])
    empty = f.slide(TITLE_AND_CONTENT)
    empty.title("Questions")  # Its content placeholder is left empty.
    closing = f.slide(TITLE_ONLY)
    closing.title("Thank you")
    closing.textbox([("Contact the operations team.", 0, "inherit")])
    closing.notes_text(["Close on time."])
    # Created welcome, agenda, internal, empty, closing; presented welcome, internal, agenda, closing, empty.
    f.reorder([0, 2, 1, 4, 3])
    return f


def bullets() -> Deck:
    """Bullets: a content placeholder's master bullets nested three levels and back, a paragraph
    whose own bullet is none between lists, a numbered list (its own auto-numbers) with a bullet
    nested in it, a bullet list straight after it at the same level, and a text box's own
    bullets and numbers beside its plain paragraphs and an empty one."""
    f = Deck("bullets")
    s = f.slide(TITLE_AND_CONTENT)
    s.title("Plan")
    s.placeholder_text(
        1,
        [
            ("Hire", 0, "inherit"),
            ("Engineers", 1, "inherit"),
            ("Backend", 2, "inherit"),
            ("Designers", 1, "inherit"),
            ("Train", 0, "inherit"),
            ("A plain paragraph between the lists.", 0, "none"),
            ("First step", 0, "number"),
            ("Second step", 0, "number"),
            ("A detail of the second step", 1, "inherit"),
            ("Third step", 0, "number"),
            ("A bullet after the numbers", 0, "inherit"),
        ],
        master_bullets,
    )
    boxes = f.slide(BLANK)
    boxes.textbox([("A text box's plain paragraph.", 0, "inherit"), ("Its own bullet", 0, "char"), ("Nested under it", 1, "char"), ("Its own number", 0, "number")])
    # An empty paragraph is spacing, not text: it is written and leaves no element.
    box = boxes.slide.shapes[-1].text_frame
    box.add_paragraph().text = ""
    box.add_paragraph().text = "After the empty paragraph."
    boxes.elements.append({"kind": "paragraph", "text": "After the empty paragraph."})
    return f


def groups() -> Deck:
    """Group shapes read in shape-tree order whatever their position: a text box, a group of a
    text box, a nested group's text box and picture, and a text box placed above them all but
    last in the tree."""
    f = Deck("groups")
    s = f.slide(BLANK)
    s.textbox([("Before the group.", 0, "inherit")], top=Inches(1))

    def outer(shapes):
        s.textbox([("In the group.", 0, "inherit")], top=Inches(2), shapes=shapes)
        inner = shapes.add_group_shape()
        s.textbox([("In the nested group.", 0, "inherit")], top=Inches(3), shapes=inner.shapes)
        s.picture(png((40, 160, 60)), "A green square in the nested group", shapes=inner.shapes)

    s.group(outer)
    s.textbox([("After the group, placed at the top.", 0, "inherit")], top=-Inches(1))
    return f


def slide_objects() -> Deck:
    """A table with merged cells across columns and rows, a cell of two paragraphs, and a pipe;
    a titled clustered column chart of two series with a missing value; the same picture on two
    slides (one file) beside another, with titles and alternative text; notes beside a chart."""
    f = Deck("slide-objects")
    red, blue = png((220, 40, 40)), png((40, 60, 200))
    t = f.slide(TITLE_ONLY)
    t.title("Figures")
    t.table(
        [
            ["Region", "Sales", None],
            ["North", "Q1\nfirst quarter", "12 | 5"],
            [None, "Q2", "14"],
        ],
        [(0, 1, 1, 2), (1, 0, 2, 1)],
    )
    c = f.slide(TITLE_ONLY)
    c.title("Sales")
    c.chart("Sales by quarter", ["Q1", "Q2", "Q3"], [("North", [12.5, 14, 15.25]), ("South", [9, None, 11])])
    c.notes_text(["The South's second quarter is missing."])
    p = f.slide(BLANK)
    p.picture(red, "A red square", "Red")
    p.picture(blue, "A blue square")
    again = f.slide(BLANK)
    again.picture(red, "The red square again")
    f.libreoffice = libreoffice_no_picture_title
    return f


PPTX_FIXTURES = [slides, bullets, groups, slide_objects]


def soffice_version() -> str:
    out = subprocess.run(["soffice", "--version"], check=True, capture_output=True, text=True)
    return out.stdout.strip()


def resave(source: pathlib.Path, target: pathlib.Path, filter: str) -> None:
    """Re-save source as target through LibreOffice headless, in a throwaway profile, with the
    export filter given."""
    with tempfile.TemporaryDirectory() as tmp:
        tmpdir = pathlib.Path(tmp)
        subprocess.run(
            [
                "soffice",
                f"-env:UserInstallation={(tmpdir / 'profile').as_uri()}",
                "--headless",
                "--convert-to",
                filter,
                "--outdir",
                str(tmpdir / "out"),
                str(source),
            ],
            check=True,
            capture_output=True,
        )
        shutil.move(tmpdir / "out" / source.name, target)


def sha256(path: pathlib.Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def update_readme(version: str, files: list[pathlib.Path]) -> None:
    begin, end = "<!-- generated: begin -->", "<!-- generated: end -->"
    text = README.read_text()
    head, rest = text.split(begin, 1)
    _, tail = rest.split(end, 1)
    lines = [f"LibreOffice: `{version}`", "", "| File | SHA-256 |", "| --- | --- |"]
    lines += [f"| `{p.relative_to(ROOT)}` | `{sha256(p)}` |" for p in files]
    README.write_text(head + begin + "\n" + "\n".join(lines) + "\n" + end + tail)


def main(names: list[str]) -> int:
    if shutil.which("soffice") is None:
        print("generate: soffice (LibreOffice) is not on PATH", file=sys.stderr)
        return 1
    # Each fixture is built once here to learn its name and its writer.
    builds, kinds = {}, {}
    for build in FIXTURES + XLSX_FIXTURES + PPTX_FIXTURES:
        fixture = build()
        builds[fixture.name] = build
        kinds[fixture.name] = (fixture.writer, fixture.extension)
    unknown = [name for name in names if name not in builds]
    if unknown:
        print(f"generate: no such fixture: {', '.join(unknown)}; the fixtures are {', '.join(builds)}", file=sys.stderr)
        return 1
    written = []
    for name in names or builds:
        fixture = builds[name]()
        out = GENERATED / fixture.extension
        out.mkdir(parents=True, exist_ok=True)
        original = out / f"{fixture.name}.{fixture.writer}.{fixture.extension}"
        resaved = out / f"{fixture.name}.libreoffice.{fixture.extension}"
        truth = out / f"{fixture.name}.truth.json"
        fixture.save(original)
        resave(original, resaved, fixture.filter)
        truth.write_text(json.dumps(fixture.truth(), indent=2, ensure_ascii=False) + "\n")
        written += [original, resaved, truth]
        variant = out / f"{fixture.name}.libreoffice.truth.json"
        adjusted = fixture.libreoffice_truth()
        if adjusted is None:
            variant.unlink(missing_ok=True)
        else:
            variant.write_text(json.dumps(adjusted, indent=2, ensure_ascii=False) + "\n")
            written.append(variant)
    # The checksums cover every fixture's files, regenerated now or not.
    files = []
    for name, (writer, extension) in kinds.items():
        suffixes = (f"{writer}.{extension}", f"libreoffice.{extension}", "truth.json", "libreoffice.truth.json")
        files += [GENERATED / extension / f"{name}.{suffix}" for suffix in suffixes]
    update_readme(soffice_version(), [path for path in files if path.exists()])
    for path in written:
        print(f"{sha256(path)}  {path.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
